"""File exports (PDF, CSV, XLSX, text). Pure Python - no Qt - so they can be unit tested headlessly."""
import csv
import json
import os
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

COPYRIGHT = 'Copyright © 2026 Nate Robertson · GPL-3.0'
_FONT_CANDIDATES = [
    ('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf',
     '/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf'),
    ('/usr/share/fonts/dejavu/DejaVuSans.ttf', '/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf',
     '/usr/share/fonts/dejavu/DejaVuSansMono.ttf'),
    ('/usr/share/fonts/TTF/DejaVuSans.ttf', '/usr/share/fonts/TTF/DejaVuSans-Bold.ttf',
     '/usr/share/fonts/TTF/DejaVuSansMono.ttf'),
]
_WIN_FONTS = Path(os.environ.get('WINDIR', r'C:\Windows')) / 'Fonts'
_FONT_CANDIDATES += [
    (str(_WIN_FONTS / 'segoeui.ttf'), str(_WIN_FONTS / 'segoeuib.ttf'), str(_WIN_FONTS / 'consola.ttf')),
    (str(_WIN_FONTS / 'arial.ttf'), str(_WIN_FONTS / 'arialbd.ttf'), str(_WIN_FONTS / 'cour.ttf')),
]


def _register_fonts():
    """Use DejaVu (Linux) or Segoe UI / Arial (Windows) for wide Unicode coverage; else built-in Helvetica/Courier."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    for regular, bold, mono in _FONT_CANDIDATES:
        if Path(regular).exists() and Path(bold).exists() and Path(mono).exists():
            try:
                pdfmetrics.registerFont(TTFont('CalderaSans', regular))
                pdfmetrics.registerFont(TTFont('CalderaSans-Bold', bold))
                pdfmetrics.registerFont(TTFont('CalderaMono', mono))
                return 'CalderaSans', 'CalderaSans-Bold', 'CalderaMono'
            except Exception:
                continue
    return 'Helvetica', 'Helvetica-Bold', 'Courier'


def build_pdf(path, title, blocks, subtitle=''):
    """blocks: list of (heading, body) tuples; heading may be ''. Body keeps its line breaks."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import HRFlowable, KeepTogether, Paragraph, SimpleDocTemplate, Spacer

    sans, bold, mono = _register_fonts()
    ink, muted, accent = colors.HexColor('#1b222c'), colors.HexColor('#667282'), colors.HexColor('#d4561f')
    s_title = ParagraphStyle('t', fontName=bold, fontSize=20, leading=25, textColor=ink, spaceAfter=2)
    s_sub = ParagraphStyle('s', fontName=sans, fontSize=9.5, leading=13, textColor=muted, spaceAfter=10)
    s_head = ParagraphStyle('h', fontName=bold, fontSize=11, leading=15, textColor=ink, spaceBefore=10, spaceAfter=3)
    s_body = ParagraphStyle('b', fontName=sans, fontSize=9.5, leading=14, textColor=ink)
    s_mono = ParagraphStyle('m', fontName=mono, fontSize=8.8, leading=12.6, textColor=ink, leftIndent=8)

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(sans, 8)
        canvas.setFillColor(muted)
        canvas.drawString(0.8 * inch, 0.5 * inch, COPYRIGHT)
        canvas.drawRightString(letter[0] - 0.8 * inch, 0.5 * inch, 'Page %d' % doc.page)
        canvas.restoreState()

    story = [Paragraph(escape(title), s_title),
             Paragraph(escape(subtitle or datetime.now().strftime('Generated %B %d, %Y  %H:%M')), s_sub),
             HRFlowable(width='100%', thickness=1.4, color=accent, spaceAfter=6)]
    for heading, body in blocks:
        group = []
        if heading:
            group.append(Paragraph(escape(str(heading)), s_head))
        text = escape(str(body).strip()).replace('\n', '<br/>')
        if text:
            group.append(Paragraph(text, s_mono if heading else s_body))
        group.append(Spacer(1, 6))
        story.append(KeepTogether(group))
    doc = SimpleDocTemplate(str(path), pagesize=letter, leftMargin=0.8 * inch, rightMargin=0.8 * inch,
                            topMargin=0.8 * inch, bottomMargin=0.85 * inch, title=title, author='Caldera')
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def project_blocks(name, notes, rows):
    blocks = []
    if str(notes).strip():
        blocks.append(('Project notes', notes))
    for r in rows:
        head = '%s   ·   %s' % (r.get('tool', ''), str(r.get('date', '')).replace('T', ' '))
        body = str(r.get('result', ''))
        if str(r.get('note', '')).strip():
            body += '\n\nNote: ' + str(r['note']).strip()
        blocks.append((head, body))
    if not rows:
        blocks.append(('', 'No calculations have been saved to this project yet.'))
    return blocks


def write_history_csv(path, rows):
    with open(path, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['Date', 'Project', 'Tool', 'Inputs', 'Result', 'Note', 'Pinned'])
        for r in rows:
            w.writerow([r.get('date'), r.get('project'), r.get('tool'), json.dumps(r.get('inputs')),
                        r.get('result'), r.get('note', ''), bool(r.get('pinned'))])


def write_history_xlsx(path, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = 'Calculation History'
    ws.append(['Date', 'Project', 'Tool', 'Inputs', 'Result', 'Note', 'Pinned'])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1B222C')
    for r in rows:
        ws.append([r.get('date'), r.get('project'), r.get('tool'), json.dumps(r.get('inputs')),
                   str(r.get('result')), str(r.get('note', '')), bool(r.get('pinned'))])
    for col, width in zip('ABCDEFG', (20, 16, 32, 40, 60, 40, 8)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def write_table_csv(path, columns, rows):
    with open(path, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(columns)
        w.writerows(rows)


REPORT_FIELDS = [('title', 'Report title'), ('asset', 'Equipment / asset'), ('tech', 'Technician'),
                 ('date', 'Test date'), ('procedure', 'Test procedure / measurements'),
                 ('results', 'Results / disposition')]


def report_text(values):
    parts = ['CALDERA TEST REPORT', '']
    for key, caption in REPORT_FIELDS:
        parts += [caption, str(values.get(key, '')).strip() or '—', '']
    return '\n'.join(parts).rstrip() + '\n'


def report_blocks(values):
    return [(caption, str(values.get(key, '')).strip() or '—') for key, caption in REPORT_FIELDS]
